import re, copy, json, sys, os
import docx
from docx.shared import Pt, Cm, Emu
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
PAGES = json.load(open(os.path.join(HERE, 'pages.json'))) if os.path.exists(os.path.join(HERE, 'pages.json')) else {}

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = nsdecls('w')
FONT = 'Cordia New'
TITLE_TH = 'การวิเคราะห์ความต้องการและแนวคิดการจัดสัดส่วนรถรับส่งตามปริมาณความต้องการนักศึกษา: กรณีศึกษาระบบรถไฟฟ้ามหาวิทยาลัยนเรศวร'
TITLE_EN = 'Demand Analysis and Fleet-Ratio Concept for the Electric Bus System: A Case Study of Naresuan University'

d = docx.Document(os.path.join(HERE, 'book.docx'))
body = d.element.body
sectPr_final = body.find(qn('w:sectPr'))

# ---------- collect front text ----------
ps = d.paragraphs
def find_h2(txt):
    for i, p in enumerate(ps):
        if p.style.name == 'Heading 2' and p.text.strip().startswith(txt):
            return i
    raise KeyError(txt)
i_abs = find_h2('บทคัดย่อ'); i_eabs = find_h2('Abstract'); i_ack = find_h2('กิตติกรรม'); i_toc = find_h2('สารบัญ')
th_abs = [ps[i].text for i in range(i_abs + 1, i_eabs)]
en_abs = [ps[i].text for i in range(i_eabs + 1, i_ack)]
ack = [ps[i].text for i in range(i_ack + 1, i_toc)]
ch1 = ps[find_h2('บทที่ 1')]._p

# ---------- delete everything before chapter 1 ----------
for el in list(body.iterchildren()):
    if el is ch1:
        break
    if el.tag == qn('w:sectPr'):
        continue
    body.remove(el)

# ---------- styles ----------
styles = d.styles.element
def set_style(sid, xml):
    old = styles.find('.//w:style[@w:styleId="%s"]' % sid, {'w': W})
    if old is not None:
        styles.remove(old)
    styles.append(parse_xml(xml))

RF = '<w:rFonts w:ascii="{f}" w:hAnsi="{f}" w:cs="{f}" w:eastAsia="{f}"/>'.format(f=FONT)
def rpr(sz, b=False, extra=''):
    return '<w:rPr>%s%s<w:sz w:val="%d"/><w:szCs w:val="%d"/>%s</w:rPr>' % (RF, '<w:b/><w:bCs/>' if b else '', sz, sz, extra)

# docDefaults
dd = styles.find(qn('w:docDefaults'))
styles.remove(dd)
styles.insert(0, parse_xml('<w:docDefaults %s><w:rPrDefault><w:rPr>%s<w:sz w:val="32"/><w:szCs w:val="32"/><w:lang w:val="en-US" w:bidi="th-TH" w:eastAsia="th-TH"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>' % (NS, RF)))

set_style('Normal', '<w:style %s w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="240" w:line="240" w:lineRule="auto"/><w:ind w:firstLine="850"/><w:jc w:val="thaiDistribute"/></w:pPr>%s</w:style>' % (NS, rpr(32)))
set_style('Heading1', '<w:style %s w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:pageBreakBefore/><w:spacing w:before="0" w:after="240"/><w:ind w:firstLine="0"/><w:jc w:val="center"/><w:outlineLvl w:val="0"/></w:pPr>%s</w:style>' % (NS, rpr(36, True)))
set_style('Heading2', '<w:style %s w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:ind w:firstLine="0"/><w:jc w:val="left"/><w:outlineLvl w:val="1"/></w:pPr>%s</w:style>' % (NS, rpr(32, True)))
set_style('Heading3', '<w:style %s w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="200" w:after="100"/><w:ind w:left="425" w:firstLine="0"/><w:jc w:val="left"/><w:outlineLvl w:val="2"/></w:pPr>%s</w:style>' % (NS, rpr(32, True)))
set_style('ListParagraph', '<w:style %s w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:after="120"/><w:ind w:left="720" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, rpr(32)))
set_style('Code', '<w:style %s w:type="paragraph" w:styleId="Code"><w:name w:val="Code"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="120"/><w:ind w:left="425" w:firstLine="0"/><w:jc w:val="left"/></w:pPr><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New"/><w:sz w:val="24"/><w:szCs w:val="24"/></w:rPr></w:style>' % NS)
set_style('Quote', '<w:style %s w:type="paragraph" w:styleId="Quote"><w:name w:val="Quote"/><w:basedOn w:val="Normal"/><w:qFormat/><w:pPr><w:ind w:left="850" w:firstLine="0"/></w:pPr>%s</w:style>' % (NS, rpr(32)))
set_style('TableCaption', '<w:style %s w:type="paragraph" w:customStyle="1" w:styleId="TableCaption"><w:name w:val="ชื่อตาราง"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="120"/><w:ind w:left="1134" w:hanging="1134"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, rpr(32)))
set_style('FigureCaption', '<w:style %s w:type="paragraph" w:customStyle="1" w:styleId="FigureCaption"><w:name w:val="ชื่อภาพ"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:spacing w:before="120" w:after="240"/><w:ind w:left="1134" w:right="0" w:hanging="1134"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, rpr(32)))
set_style('FrontHeading', '<w:style %s w:type="paragraph" w:customStyle="1" w:styleId="FrontHeading"><w:name w:val="หัวข้อส่วนนำ"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:qFormat/><w:pPr><w:keepNext/><w:spacing w:before="0" w:after="240"/><w:ind w:firstLine="0"/><w:jc w:val="center"/></w:pPr>%s</w:style>' % (NS, rpr(36, True)))
set_style('TableNote', '<w:style %s w:type="paragraph" w:customStyle="1" w:styleId="TableNote"><w:name w:val="หมายเหตุตาราง"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:before="60" w:after="240"/><w:ind w:firstLine="0"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, rpr(28)))
TABW = 8362
for lvl, sid in ((1, 'TOC1'), (2, 'TOC2'), (3, 'TOC3')):
    set_style(sid, '<w:style %s w:type="paragraph" w:styleId="%s"><w:name w:val="toc %d"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="%d"/></w:tabs><w:spacing w:after="0"/><w:ind w:left="%d" w:hanging="0" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, sid, lvl, TABW, (lvl - 1) * 425, rpr(32, lvl == 1)))
set_style('TableofFigures', '<w:style %s w:type="paragraph" w:styleId="TableofFigures"><w:name w:val="table of figures"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="%d"/></w:tabs><w:spacing w:after="0"/><w:ind w:left="1134" w:hanging="1134" w:firstLine="0"/><w:jc w:val="left"/></w:pPr>%s</w:style>' % (NS, TABW, rpr(32)))

# ---------- text edits on the kept body ----------
EDITS = [
    (re.compile(r'ตารางที่ (?=[0-9ก-ค]+\.[0-9])'), 'ตาราง '),
    (re.compile(r'ภาพที่ (?=[0-9]+\.[0-9])'), 'ภาพ '),
    (re.compile(r'Campus Shuttle System'), 'Electric Bus System'),
    ('คร่าว ๆ', 'โดยประมาณ'), ('ตรง ๆ', 'โดยตรง'), ('สั้น ๆ', 'สั้น'),
    ('เพื่อน ๆ', 'เพื่อน'),
]
def edit_text(root):
    for t in root.iter(qn('w:t')):
        s = t.text or ''
        o = s
        for pat, rep in EDITS:
            s = pat.sub(rep, s) if hasattr(pat, 'sub') else s.replace(pat, rep)
        if s != o:
            t.text = s
edit_text(body)

# ---------- helpers to build elements ----------
def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def run_xml(text, sz=32, b=False, font=FONT, i=False):
    f = '<w:rFonts w:ascii="{f}" w:hAnsi="{f}" w:cs="{f}" w:eastAsia="{f}"/>'.format(f=font)
    return '<w:r><w:rPr>%s%s%s<w:sz w:val="%d"/><w:szCs w:val="%d"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (f, '<w:b/><w:bCs/>' if b else '', '<w:i/><w:iCs/>' if i else '', sz, sz, esc(text))

def para_xml(runs, jc='left', first=0, after=0, before=0, style=None, keep=False, left=0, extra='', pagebreak=False):
    ppr = ''
    if style:
        ppr += '<w:pStyle w:val="%s"/>' % style
    if keep:
        ppr += '<w:keepNext/>'
    if pagebreak:
        ppr += '<w:pageBreakBefore/>'
    ppr += extra
    ppr += '<w:spacing w:before="%d" w:after="%d"/><w:ind w:left="%d" %s/><w:jc w:val="%s"/>' % (before, after, left, ('w:hanging="%d"' % -first) if first < 0 else ('w:firstLine="%d"' % first), jc)
    return '<w:p %s><w:pPr>%s</w:pPr>%s</w:p>' % (NS, ppr, runs)

front = []  # list of lxml elements in order
def add(xml):
    front.append(parse_xml(xml))

def sect_xml(left, right, top, bottom, header=1418, fmt=None, start=None, brk='nextPage', headref=None):
    pg = ''
    if fmt or start is not None:
        pg = '<w:pgNumType%s%s/>' % (' w:fmt="%s"' % fmt if fmt else '', ' w:start="%d"' % start if start is not None else '')
    return ('<w:sectPr %s><w:type w:val="%s"/><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="%d" w:right="%d" w:bottom="%d" w:left="%d" w:header="%d" w:footer="709" w:gutter="0"/>%s</w:sectPr>'
            % (NS, brk, top, right, bottom, left, header, pg))

def sect_para(sect):
    return parse_xml('<w:p %s><w:pPr>%s</w:pPr></w:p>' % (NS, sect.replace(' ' + NS, '', 1) if False else sect.split('>', 1)[0] and ''))

def sect_par(left, right, top, bottom, **kw):
    s = sect_xml(left, right, top, bottom, **kw)
    # wrap sectPr (drop its own namespace decl) in an empty paragraph
    s2 = s.replace(' ' + NS, '', 1)
    return parse_xml('<w:p %s><w:pPr><w:spacing w:after="0"/>%s</w:pPr></w:p>' % (NS, s2))

# ---------- cover (section 1) ----------
LOGO = os.path.join(HERE, 'x8', 'word', 'media', 'image1.png')
cp = d.add_paragraph()
cp.alignment = 1
cp.paragraph_format.first_line_indent = 0
cp.add_run().add_picture(LOGO, width=Cm(3.5))
logo_el = cp._p
body.remove(logo_el)
front.append(logo_el)
add(para_xml(run_xml(TITLE_TH, 40, True, 'TH SarabunPSK'), 'center', after=240, before=300))
add(para_xml(run_xml(TITLE_EN, 40, True, 'TH SarabunPSK'), 'center', after=800))
add(para_xml(run_xml('กรกช หิริรัตน์', 40, True, 'TH SarabunPSK'), 'center', after=700))
for line in ['วิทยานิพนธ์ระดับปริญญาตรี เสนอภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม',
             'เพื่อเป็นส่วนหนึ่งของการศึกษาหลักสูตรปริญญาวิทยาศาสตรบัณฑิต',
             'สาขาวิชาภูมิศาสตร์', 'ตุลาคม 2569', 'ลิขสิทธิ์เป็นของมหาวิทยาลัยนเรศวร']:
    add(para_xml(run_xml(line, 40, True, 'TH SarabunPSK'), 'center', after=60))
front.append(sect_par(1440, 1411, 2131, 1411))

# ---------- approval / info page (section 2) ----------
def kv_table(rows, w1=2900, w2=5462, sz=32, b1=True):
    trs = ''
    for k, v in rows:
        trs += '<w:tr><w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/></w:tcPr>%s</w:tc><w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/></w:tcPr>%s</w:tc></w:tr>' % (
            w1, para_xml(run_xml(k, sz, b1), 'left', after=120), w2, para_xml(run_xml(v, sz), 'left', after=120))
    return parse_xml('<w:tbl %s><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid><w:gridCol w:w="%d"/><w:gridCol w:w="%d"/></w:tblGrid>%s</w:tbl>' % (NS, w1 + w2, w1, w2, trs))

logo2 = d.add_paragraph(); logo2.alignment = 1; logo2.paragraph_format.first_line_indent = 0
logo2.add_run().add_picture(LOGO, width=Cm(3.0)); logo2.paragraph_format.space_after = Pt(12)
body.remove(logo2._p); front.append(logo2._p)
front.append(kv_table([
    ('ชื่อเรื่องวิทยานิพนธ์ภาษาไทย', TITLE_TH),
    ('ชื่อเรื่องวิทยานิพนธ์ภาษาอังกฤษ', TITLE_EN),
    ('ผู้ศึกษาวิจัย', 'นางสาวกรกช หิริรัตน์'),
    ('ที่ปรึกษาวิทยานิพนธ์', 'ผู้ช่วยศาสตราจารย์ ดร.กัมปนาท ปิยะธำรงชัย'),
    ('ประเภทสารนิพนธ์', 'วิทยานิพนธ์ระดับปริญญาตรี วท.บ. ภูมิศาสตร์ มหาวิทยาลัยนเรศวร 2569')]))
add(para_xml(run_xml('อาจารย์ที่ปรึกษา ประธานหลักสูตรวิทยาศาสตรบัณฑิต สาขาวิชาภูมิศาสตร์ และหัวหน้าภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม ได้พิจารณาวิทยานิพนธ์ เรื่อง “' + TITLE_TH + '” (' + TITLE_EN + ') ของ กรกช หิริรัตน์ เห็นสมควรรับเป็นส่วนหนึ่งของการศึกษาตามหลักสูตรปริญญาวิทยาศาสตรบัณฑิต สาขาวิชาภูมิศาสตร์ ของมหาวิทยาลัยนเรศวร'), 'thaiDistribute', first=850, before=360, after=480))
sigs = [('(ผู้ช่วยศาสตราจารย์ ดร.กัมปนาท ปิยะธำรงชัย)', 'อาจารย์ที่ปรึกษาวิทยานิพนธ์'),
        ('(อาจารย์ ธัญญาลัคน์ จันทร์สมบัติ)', 'ประธานหลักสูตรปริญญาวิทยาศาสตรบัณฑิต สาขาวิชาภูมิศาสตร์'),
        ('(ผู้ช่วยศาสตราจารย์ ว่าที่ ร.ต.ดร.รังสรรค์ เกตุออ)', 'หัวหน้าภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม')]
cw = TABW // 3
cells = ''
for nm, role in sigs:
    cells += '<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/></w:tcPr>%s%s%s</w:tc>' % (
        cw, para_xml(run_xml('…………………………………', 28), 'center', before=600, after=60),
        para_xml(run_xml(nm, 28), 'center', after=60), para_xml(run_xml(role, 28), 'center', after=60))
front.append(parse_xml('<w:tbl %s><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>%s</w:tblGrid><w:tr>%s</w:tr></w:tbl>' % (NS, cw * 3, '<w:gridCol w:w="%d"/>' % cw * 3, cells)))
add(para_xml('', 'left'))
front.append(sect_par(2126, 1418, 2126, 1418))

# ---------- English abstract page (section 3: no number) ----------
add(para_xml(run_xml('Title', 32, True) + run_xml('\t' + TITLE_EN, 32), 'left', after=60, left=1700, extra='', first=-1700))
kv = [('Title', TITLE_EN), ('Author', 'Korakot Hirirat'), ('Advisor', 'Assistant Professor Kampanat Piyathamrongchai, Ph.D.'),
      ('Academic Paper', 'Undergraduate Thesis B.Sc. Geography, Naresuan University, 2026'),
      ('Keywords', en_abs[-1].replace('Keywords:', '').strip())]
front.pop()
front.append(kv_table(kv, 2200, 6162))
add(para_xml(run_xml('ABSTRACT', 36, True), 'center', before=360, after=240))
for t in en_abs[:-1]:
    t = t.replace('Campus Shuttle System', 'Electric Bus System')
    add(para_xml(run_xml(t), 'both', first=850, after=120))
front.append(sect_par(2126, 1418, 2126, 1418))

# ---------- Thai abstract + acknowledgements + lists (section 4: Thai-letter numbers) ----------
add(para_xml(run_xml('บทคัดย่อ', 36, True), 'center', after=240))
for t in th_abs[:-1]:
    add(para_xml(run_xml(t), 'thaiDistribute', first=850, after=120))
kw = th_abs[-1]
if ':' in kw:
    a, b = kw.split(':', 1)
    add(para_xml(run_xml(a + ':', 32, True) + run_xml(b), 'left', after=120, before=120))
add(para_xml(run_xml('ประกาศคุณูปการ', 36, True), 'center', after=240, pagebreak=True))
for t in ack[:-1]:
    add(para_xml(run_xml(t), 'thaiDistribute', first=850, after=240))
add(para_xml(run_xml(ack[-1]), 'right', before=480))

def field_para(instr, cached, style, pagebreak=False):
    return parse_xml('<w:p %s><w:pPr><w:pStyle w:val="%s"/></w:pPr><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> %s </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>%s</w:p>' % (NS, style, instr, cached))

HEADS = []  # filled after parse: (level, text) for TOC cache

def collect_headings():
    out = []
    for el in body.iterchildren():
        if el.tag != qn('w:p'):
            continue
        p = Paragraph(el, d)
        sn = p.style.name
        if sn in ('Heading 2', 'Heading 3', 'Heading 4'):
            out.append(sn)
    return out

def toc_entry(text, page, style):
    return '<w:r><w:t xml:space="preserve">%s</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>%s</w:t></w:r>' % (esc(text), page)

# ---------- chapter / heading restyle in body ----------
for el in list(body.iterchildren()):
    if el.tag != qn('w:p'):
        continue
    p = Paragraph(el, d)
    sn = p.style.name
    txt = p.text.strip()
    if sn == 'Heading 2':
        m = re.match(r'^(บทที่ \d+|ภาคผนวก [ก-ฮ])\s+(.*)$', txt)
        for r in list(el.findall(qn('w:r'))):
            el.remove(r)
        for ch in list(el):
            if ch.tag != qn('w:pPr'):
                el.remove(ch)
        p.style = d.styles['Heading 1']
        if m:
            el.append(parse_xml('<w:r %s>%s<w:br/><w:t xml:space="preserve">%s</w:t></w:r>' % (NS, '<w:t>%s</w:t>' % esc(m.group(1)), esc(m.group(2)))))
        else:
            el.append(parse_xml('<w:r %s><w:t xml:space="preserve">%s</w:t></w:r>' % (NS, esc(txt))))
    elif sn == 'Heading 3':
        p.style = d.styles['Heading 2']
    elif sn == 'Heading 4':
        p.style = d.styles['Heading 3']
    elif sn == 'List Paragraph':
        pass

# captions + notes + tables + pictures
bl = list(body.iterchildren())
for i, el in enumerate(bl):
    if el.tag == qn('w:p'):
        t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
        p = Paragraph(el, d)
        has_pic = len(list(el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))) > 0
        if has_pic:
            p.alignment = 1
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            for dr in el.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/wordprocessingDrawing}extent'):
                cx, cy = int(dr.get('cx')), int(dr.get('cy'))
                mx = int(Cm(14.6))
                if cx > mx:
                    dr.set('cx', str(mx)); dr.set('cy', str(int(cy * mx / cx)))
            for ex in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext'):
                if ex.get('cx'):
                    cx, cy = int(ex.get('cx')), int(ex.get('cy'))
                    mx = int(Cm(14.6))
                    if cx > mx:
                        ex.set('cx', str(mx)); ex.set('cy', str(int(cy * mx / cx)))
            continue
        nxt = bl[i + 1] if i + 1 < len(bl) else None
        prv = bl[i - 1] if i > 0 else None
        if re.match(r'^ตาราง [0-9ก-ค]+\.[0-9]+ ', t) and nxt is not None and nxt.tag == qn('w:tbl'):
            p.style = d.styles['ชื่อตาราง']
        elif re.match(r'^ภาพ [0-9]+\.[0-9]+ ', t) and prv is not None and len(list(prv.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))):
            p.style = d.styles['ชื่อภาพ']
        elif re.match(r'^(หมายเหตุ|ที่มา)', t) and prv is not None and (prv.tag == qn('w:tbl') or prv.tag == qn('w:p')):
            p.style = d.styles['หมายเหตุตาราง']

# direct run-level fonts: strip anything that might override
for rp in body.iter(qn('w:rPr')):
    for tag in ('w:rFonts', 'w:sz', 'w:szCs'):
        for x in rp.findall(qn(tag)):
            rp.remove(x)

# tables
B = '<w:{n} w:val="single" w:sz="{s}" w:space="0" w:color="000000"/>'
for tbl in body.iter(qn('w:tbl')):
    if tbl.getparent() is not body:
        continue
    ncols = len(tbl.find(qn('w:tblGrid')).findall(qn('w:gridCol')))
    tp = tbl.find(qn('w:tblPr'))
    for ch in list(tp):
        tp.remove(ch)
    tp.append(parse_xml('<w:tblW %s w:w="5000" w:type="pct"/>' % NS))
    tp.append(parse_xml('<w:jc %s w:val="center"/>' % NS))
    tp.append(parse_xml('<w:tblBorders %s>%s%s%s</w:tblBorders>' % (NS, B.format(n='top', s=12), B.format(n='bottom', s=12), B.format(n='insideH', s=4))))
    tp.append(parse_xml('<w:tblLayout %s w:type="autofit"/>' % NS))
    tp.append(parse_xml('<w:tblCellMar %s><w:left w:w="57" w:type="dxa"/><w:right w:w="57" w:type="dxa"/></w:tblCellMar>' % NS))
    sz = 28 if ncols <= 5 else (24 if ncols <= 9 else 20)
    grid = tbl.find(qn('w:tblGrid'))
    rows = tbl.findall(qn('w:tr'))
    for ri, tr in enumerate(rows):
        trp = tr.find(qn('w:trPr'))
        if trp is None:
            trp = parse_xml('<w:trPr %s/>' % NS); tr.insert(0, trp)
        trp.append(parse_xml('<w:cantSplit %s/>' % NS))
        if ri == 0:
            trp.append(parse_xml('<w:tblHeader %s/>' % NS))
        for tc in tr.findall(qn('w:tc')):
            tcp = tc.find(qn('w:tcPr'))
            if tcp is None:
                tcp = parse_xml('<w:tcPr %s/>' % NS); tc.insert(0, tcp)
            for ch in list(tcp):
                if ch.tag in (qn('w:shd'), qn('w:tcBorders'), qn('w:tcW')):
                    tcp.remove(ch)
            if ri == 0:
                tcp.append(parse_xml('<w:tcBorders %s><w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/></w:tcBorders>' % NS))
            for pp in tc.findall(qn('w:p')):
                ppr = pp.find(qn('w:pPr'))
                if ppr is None:
                    ppr = parse_xml('<w:pPr %s/>' % NS); pp.insert(0, ppr)
                for ch in list(ppr):
                    ppr.remove(ch)
                ppr.append(parse_xml('<w:spacing %s w:before="20" w:after="20"/>' % NS))
                ppr.append(parse_xml('<w:ind %s w:left="0" w:firstLine="0"/>' % NS))
                ppr.append(parse_xml('<w:jc %s w:val="left"/>' % NS))
                for r in pp.iter(qn('w:r')):
                    rp = r.find(qn('w:rPr'))
                    if rp is None:
                        rp = parse_xml('<w:rPr %s/>' % NS); r.insert(0, rp)
                    if ri == 0 and rp.find(qn('w:b')) is None:
                        rp.insert(0, parse_xml('<w:b %s/>' % NS))
                    rp.append(parse_xml('<w:sz %s w:val="%d"/>' % (NS, sz)))
                    rp.append(parse_xml('<w:szCs %s w:val="%d"/>' % (NS, sz)))

# ---------- biography ----------
bio_rows = [
    ('ชื่อ – ชื่อสกุล', 'กรกช หิริรัตน์'),
    ('วัน เดือน ปี เกิด', '3 พฤษภาคม 2547'),
    ('ที่อยู่ปัจจุบัน', '165/1 หมู่ 4 ตำบลงิ้วราย อำเภอตะพานหิน จังหวัดพิจิตร 66110'),
]
bio_edu = [('พ.ศ. 2566–ปัจจุบัน', 'ระดับปริญญาตรี หลักสูตรวิทยาศาสตรบัณฑิต สาขาวิชาภูมิศาสตร์ มหาวิทยาลัยนเรศวร ตำบลท่าโพธิ์ อำเภอเมืองพิษณุโลก จังหวัดพิษณุโลก 65000 เกรดเฉลี่ย 3.75'),
           ('พ.ศ. 2562–2564', 'ระดับมัธยมศึกษาตอนปลาย (วิทย์–คณิต) โรงเรียนศรีราชา ตำบลศรีราชา อำเภอศรีราชา จังหวัดชลบุรี'),
           ('พ.ศ. 2559–2561', 'ระดับมัธยมศึกษาตอนต้น โรงเรียนตะพานหิน ตำบลตะพานหิน อำเภอตะพานหิน จังหวัดพิจิตร')]
bio_act = ['ผู้ช่วยสอน (Teacher Assistant – TA) รายวิชาแอนโทรโปซีน 2 ภาคการศึกษา [2567–2568]',
           'ดำรงตำแหน่งประธานชมรมภูมิศาสตร์ พ.ศ. 2568',
           'ประธานกิจกรรม โครงการการพัฒนาทักษะการเรียนรู้ตลอดชีวิต (ค่ายสานสัมพันธ์พี่น้องประจำปี 2568 และ 2569 กิจกรรมจัดขึ้น ณ อุทยานแห่งชาติภูหินร่องกล้า)',
           'สตาฟกิจกรรมภูมิศาสตร์ สอวน. ค่าย 1–2 ประจำปี 2568 มหาวิทยาลัยนเรศวร']
bio_aw = ['ได้รับเกียรติบัตรด้านการเรียนผลการเรียนดี ประจำปีการศึกษา 2566 สาขาวิชาภูมิศาสตร์ ภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม มหาวิทยาลัยนเรศวร',
          'ได้รับเกียรติบัตรด้านการเรียนผลการเรียนดี ประจำปีการศึกษา 2567 สาขาวิชาภูมิศาสตร์ ภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม มหาวิทยาลัยนเรศวร',
          'ได้รับเกียรติบัตรด้านการเรียนผลการเรียนดี ประจำปีการศึกษา 2568 สาขาวิชาภูมิศาสตร์ ภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม มหาวิทยาลัยนเรศวร',
          'ได้รับเกียรติบัตรด้านพัฒนาการเรียนยอดเยี่ยม ประจำปีการศึกษา 2568 สาขาวิชาภูมิศาสตร์ ภาควิชาทรัพยากรธรรมชาติและสิ่งแวดล้อม มหาวิทยาลัยนเรศวร',
          'ได้รับเกียรติบัตรนิสิตผลการเรียนดี ประจำปีการศึกษา 2567 โดยมหาวิทยาลัยนเรศวร',
          'ได้รับใบประกาศเกียรติคุณ จากการประกวดโมเดลธุรกิจ ด้านส่งเสริมชุมชนและตลาดสร้างสรรค์ ในฐานะผู้จัดการโครงการ “Little Garden” ภายใต้โครงการเรียนรู้และฝึกประสบการณ์นิสิต รุ่นที่ 1']

tail = []
def tadd(xml):
    tail.append(parse_xml(xml))
# remove old bio section (heading 'ประวัติผู้วิจัย' to end)
els = list(body.iterchildren())
start = None
for i, el in enumerate(els):
    if el.tag == qn('w:p') and ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip() == 'ประวัติผู้วิจัย':
        start = i
if start is not None:
    for el in els[start:]:
        if el.tag != qn('w:sectPr'):
            body.remove(el)
tadd('<w:p %s><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>ประวัติผู้วิจัย</w:t></w:r></w:p>' % NS)
tail.append(kv_table(bio_rows, 2800, 5562))
tadd(para_xml(run_xml('ประวัติการศึกษา', 32, True), 'left', before=240, after=60, keep=True))
tail.append(kv_table(bio_edu, 2300, 6062, b1=False))
tadd(para_xml(run_xml('กิจกรรมที่เข้าร่วม', 32, True), 'left', before=240, after=60, keep=True))
for a in bio_act:
    tadd(para_xml(run_xml('– ' + a), 'left', left=425, first=-425, after=60))
tadd(para_xml(run_xml('รางวัลที่ได้รับ', 32, True), 'left', before=240, after=60, keep=True))
for a in bio_aw:
    tadd(para_xml(run_xml('– ' + a), 'left', left=425, first=-425, after=60))

# ---------- headings cache for TOC ----------
def heading_list():
    out = []
    for el in body.iterchildren():
        if el.tag != qn('w:p'):
            continue
        p = Paragraph(el, d)
        sn = p.style.name
        if sn in ('Heading 1', 'Heading 2'):
            txt = ''.join((' ' if x.tag == qn('w:br') else (x.text or '')) for x in el.iter() if x.tag in (qn('w:t'), qn('w:br')))
            out.append((1 if sn == 'Heading 1' else 2, re.sub(r'\s+', ' ', txt).strip()))
    return out
for t in tail:
    body.insert(list(body).index(sectPr_final), t)
HL = heading_list()
CAP_T = []
CAP_F = []
for el in body.iterchildren():
    if el.tag == qn('w:p'):
        p = Paragraph(el, d)
        if p.style.name == 'ชื่อตาราง':
            CAP_T.append(re.sub(r'\s+', ' ', p.text).strip())
        elif p.style.name == 'ชื่อภาพ':
            CAP_F.append(re.sub(r'\s+', ' ', p.text).strip())

def entries(items, style, key):
    out = []
    for lvl, text in items:
        pg = PAGES.get(key + '|' + text, '')
        st = style if isinstance(style, str) else style[lvl]
        out.append('<w:p %s><w:pPr><w:pStyle w:val="%s"/></w:pPr>%s</w:p>' % (NS, st, toc_entry(text, pg, st)))
    return out

def field_block(instr, items, style, key, title, headerline):
    out = [parse_xml(para_xml(run_xml(title, 36, True), 'center', after=240, pagebreak=True))]
    out.append(parse_xml(para_xml(run_xml(headerline[0], 32, True) + '<w:r><w:tab/></w:r>' + run_xml(headerline[1], 32, True), 'left', after=60,
                                  extra='<w:tabs><w:tab w:val="right" w:pos="%d"/></w:tabs>' % TABW)))
    es = entries(items, style, key)
    for k, e in enumerate(es):
        x = e
        if k == 0:
            x = x.replace('</w:pPr>', '</w:pPr><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> %s </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>' % instr, 1)
        if k == len(es) - 1:
            x = x.replace('</w:p>', '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
        out.append(parse_xml(x))
    return out

front += field_block('TOC \\o &quot;1-2&quot; \\h \\z \\u', HL, {1: 'TOC1', 2: 'TOC2'}, 'h', 'สารบัญ', ('บทที่', 'หน้า'))
front += field_block('TOC \\h \\z \\t &quot;ชื่อตาราง,1&quot;', [(1, t) for t in CAP_T], 'TableofFigures', 'h', 'สารบัญตาราง', ('ตาราง', 'หน้า'))
front += field_block('TOC \\h \\z \\t &quot;ชื่อภาพ,1&quot;', [(1, t) for t in CAP_F], 'TableofFigures', 'h', 'สารบัญภาพ', ('ภาพ', 'หน้า'))
front.append(sect_par(2126, 1418, 2126, 1418, fmt='thaiLetters', start=1))

# ---------- insert front before chapter 1 ----------
for el in front:
    ch1.addprevious(el)

# ---------- final section properties ----------
for ch in list(sectPr_final):
    sectPr_final.remove(ch)
for ch in parse_xml(sect_xml(2126, 1418, 2126, 1418, fmt='decimal', start=1)):
    sectPr_final.append(ch)

# ---------- headers with page numbers ----------
def add_header(section):
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp._p.append(parse_xml('<w:pPr %s><w:jc w:val="right"/></w:pPr>' % NS)) if hp._p.find(qn('w:pPr')) is None else None
    ppr = hp._p.find(qn('w:pPr'))
    for ch in list(ppr):
        ppr.remove(ch)
    ppr.append(parse_xml('<w:ind %s w:firstLine="0"/>' % NS))
    ppr.append(parse_xml('<w:jc %s w:val="right"/>' % NS))
    for x in ('<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>', '<w:r %s><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>', '<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>', '<w:r %s><w:t>1</w:t></w:r>', '<w:r %s><w:fldChar w:fldCharType="end"/></w:r>'):
        r = parse_xml(x % NS)
        rp = parse_xml('<w:rPr %s>%s<w:sz w:val="32"/><w:szCs w:val="32"/></w:rPr>' % (NS, RF))
        r.insert(0, rp)
        hp._p.append(r)

secs = d.sections
# sections: 0 cover, 1 approval, 2 english abstract, 3 front matter, 4 body
for k, s in enumerate(secs):
    if k >= 3:
        add_header(s)
    else:
        s.header.is_linked_to_previous = False

# settings: update fields on open
st = d.settings.element
uf = parse_xml('<w:updateFields %s w:val="true"/>' % NS)
after_tags = ['hdrShapeDefaults','footnotePr','endnotePr','compat','docVars','rsids','mathPr','attachedSchema','themeFontLang','clrSchemeMapping','doNotIncludeSubdocsInStats','doNotAutoCompressPictures','forceUpgrade','captions','readModeInkLockDown','smartTagType','schemaLibrary','shapeDefaults','doNotEmbedSmartTags','decimalSymbol','listSeparator']
for ch in st:
    if ch.tag.split('}')[1] in after_tags:
        ch.addprevious(uf); break
else:
    st.append(uf)

d.core_properties.title = TITLE_TH
d.core_properties.author = 'กรกช หิริรัตน์'
d.save(os.path.join(HERE, 'thesis_geo.docx'))
print('saved', len(secs), 'sections', len(HL), 'heads', len(CAP_T), 'tables', len(CAP_F), 'figs')
